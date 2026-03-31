from __future__ import annotations

import os
from datetime import date
from typing import Iterable

from docx import Document
from docx.shared import Pt

from config import (
    TEMPLATE_PATH, OUTPUT_DIR, OUTPUT_NAME_FMT, NEWS_COUNT,
    FONT_CN_BODY, FONT_EN, FONT_SIZE_PT_SANHAO
)
from core.model import Report
from core.placeholders import ISSUE_NO, DATE_Y, DATE_M, DATE_D, news_placeholders
from core.text_runs import split_cn_en


def _iter_all_paragraphs(doc: Document) -> Iterable:
    """遍历正文段落 + 表格内段落（模板一般只用正文段落，但这更稳）"""
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def _clear_runs(paragraph) -> None:
    """清空一个段落的所有 run，但保留段落格式（缩进/行距/对齐等）"""
    for r in paragraph.runs:
        r._element.getparent().remove(r._element)


def _add_mixed_runs(paragraph, text: str, bold: bool) -> None:
    """按中英数字拆分写入 run：中文仿宋，英数 Times New Roman"""
    segments = split_cn_en(text)
    for lang, seg in segments:
        run = paragraph.add_run(seg)
        run.bold = bold
        run.font.size = Pt(FONT_SIZE_PT_SANHAO)
        run.font.name = FONT_CN_BODY if lang == "cn" else FONT_EN


def _replace_simple_placeholders(doc: Document, mapping: dict[str, str]) -> None:
    """
    替换类似 ISSUE_NO / DATE_Y 的“短占位符”：
    只做 run 内替换，不改变 run 样式
    """
    for p in _iter_all_paragraphs(doc):
        for r in p.runs:
            for k, v in mapping.items():
                if k in r.text:
                    r.text = r.text.replace(k, v)


def export_report(report: Report) -> str:
    """
    读取模板 -> 替换占位符 -> run级重写标题/正文 -> 输出docx
    返回输出文件路径
    """
    if len(report.items) != NEWS_COUNT:
        raise ValueError(f"items must be {NEWS_COUNT}, got {len(report.items)}")

    doc = Document(TEMPLATE_PATH)

    # 1) 顶部简单占位符：期号、日期（按 run 内替换，保持模板字体）
    simple_map = {
        ISSUE_NO: str(report.issue_no),
        DATE_Y: str(report.report_date.year),
        DATE_M: str(report.report_date.month),
        DATE_D: str(report.report_date.day),
    }
    _replace_simple_placeholders(doc, simple_map)

    # 2) 逐条处理：找到包含 {{N1_TITLE}} / {{N1_BODY}} 的段落，整段重写
    paragraphs = list(_iter_all_paragraphs(doc))

    for i, item in enumerate(report.items, start=1):
        ph = news_placeholders(i)

        # 标题整段重写（整段加粗）
        title_text = (
            f"{i}.【{item.topic}】{item.title}"
            f"（{item.media} {item.source_date.year}-{item.source_date.month}-{item.source_date.day}）"
        )

        # 正文整段重写（不加粗）
        body_text = item.body

        # 找标题段落
        title_p = next((p for p in paragraphs if ph.title in p.text), None)
        if title_p is None:
            raise RuntimeError(f"Cannot find title placeholder for item {i}: {ph.title}")
        _clear_runs(title_p)
        _add_mixed_runs(title_p, title_text, bold=True)

        # 找正文段落
        body_p = next((p for p in paragraphs if ph.body in p.text), None)
        if body_p is None:
            raise RuntimeError(f"Cannot find body placeholder for item {i}: {ph.body}")
        _clear_runs(body_p)
        _add_mixed_runs(body_p, body_text, bold=False)

    # 3) 输出
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    yyyymmdd = report.report_date.strftime("%Y%m%d")
    out_name = OUTPUT_NAME_FMT.format(yyyymmdd=yyyymmdd)
    out_path = os.path.join(OUTPUT_DIR, out_name)
    doc.save(out_path)
    return out_path